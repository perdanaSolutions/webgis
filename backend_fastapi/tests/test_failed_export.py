"""Unduhan Excel data gagal tahap SIMPAN (karantina) + gagal_lain di table_preview (tanpa database)."""
import io
import uuid
from datetime import datetime, timezone

import pytest
from openpyxl import load_workbook

from app.services.upload import failed_export, part_response
from app.services.upload.failed_export import CELL_LIMIT, TRUNC_NOTE


def _feature(no, **props):
    return {"type": "Feature", "properties": {"OBJECTID": no, **props},
            "geometry": {"type": "Point", "coordinates": [110.0 + no / 1000, 0.5]}}


def _info(rows, total=None):
    return {"kind": "execute", "layer": "blok", "filename": "kebun.geojson", "bulan": 1, "tahun": 2026,
            "waktu": datetime(2026, 10, 27, 1, 0, tzinfo=timezone.utc), "total": total, "part_total": 3,
            "rows": sorted(rows, key=failed_export._sort_key)}


def _sheets(content):
    wb = load_workbook(io.BytesIO(content))
    return wb["Data Gagal"], wb["Ringkasan"]


def test_table_preview_keeps_all_reasons_per_row():
    rows = part_response.table_preview(3, [
        {"no_preview": 2, "kode": "GEOMETRI_TIDAK_VALID", "alasan": "a"},
        {"no_preview": 2, "kode": "BLOK_TIDAK_DITEMUKAN", "alasan": "lihat data #3"},
    ], offset=1000)
    assert rows[1]["kode_gagal"] == "GEOMETRI_TIDAK_VALID"
    assert rows[1]["gagal_lain"] == [{"kode": "BLOK_TIDAK_DITEMUKAN", "alasan": "lihat data #1003"}]
    assert "gagal_lain" not in rows[0] and "gagal_lain" not in rows[2]


def test_workbook_contains_all_quarantine_rows_without_cap():
    rows = [{"no": n, "kode": "GEOMETRI_TIDAK_VALID", "alasan": f"rusak {n}", "sumber": _feature(n, Blok=f"B{n}")}
            for n in range(1, 3501)]
    rows.append({"no": 4000, "kode": "GAGAL_DISIMPAN", "alasan": "x", "sumber": _feature(4000, EstID="E001")})
    data, summary = _sheets(failed_export.build_workbook(_info(rows)))
    assert data.max_row == 3502
    header = [c.value for c in data[1]]
    assert header[:9] == failed_export.FIXED_HEADERS + ["Estate", "Afdeling", "Blok", "OBJECTID"]
    last = [c.value for c in data[3502]]
    assert last[1] == 4000 and last[2] == "GAGAL_DISIMPAN" and last[5] == "E001" and last[-1].startswith("POINT")
    assert data.freeze_panes == "A2"
    counts = {r[0].value: r[1].value for r in summary.iter_rows(min_row=12) if r[0].value}
    assert counts == {"GEOMETRI_TIDAK_VALID": 3500, "GAGAL_DISIMPAN": 1}


@pytest.mark.parametrize("value", ["=HYPERLINK(\"x\")", "+1", "-2+3", "@SUM(A1)", "\t=1"])
def test_formula_injection_is_escaped(value):
    assert failed_export.safe_cell(value) == "'" + value


def test_long_cell_is_truncated_with_note():
    out = failed_export.safe_cell("x" * 40_000)
    assert len(out) == CELL_LIMIT and out.endswith(TRUNC_NOTE)


def test_excel_rows_without_geometry_column():
    data, _ = _sheets(failed_export.build_workbook(_info(
        [{"no": 1, "kode": "NILAI_TIDAK_VALID", "alasan": "=bad", "sumber": {"Blok": "B1", "Luas": 2}}])))
    assert "Geometri (WKT)" not in [c.value for c in data[1]] and data.cell(2, 5).value == "'=bad"


def test_filename_uses_wib_and_clean_name():
    info = {"layer": "blok", "filename": "Kebun A (final)_part1of3.geojson"}
    now = datetime(2026, 10, 27, 0, 5, tzinfo=timezone.utc)
    assert failed_export.export_filename(info, now) == "data_gagal_blok_Kebun_A_final_20261027-0705.xlsx"


def test_endpoint_requires_auth():
    from fastapi.testclient import TestClient
    from app.main import app
    r = TestClient(app).get(f"/api/v1/upload-jobs/{uuid.uuid4()}/data-gagal.xlsx")
    assert r.status_code in (401, 403)


def test_endpoint_returns_xlsx_and_404_when_empty(monkeypatch):
    from types import SimpleNamespace
    from fastapi.testclient import TestClient
    from app.api.deps import get_current_user, get_db
    from app.api.v1.endpoints import upload_jobs as ep
    from app.main import app

    info = _info([{"no": 1, "kode": "OBJECTID_GANDA", "alasan": "dobel", "sumber": _feature(1)}])
    monkeypatch.setattr(ep.failed_export, "collect", lambda db, job_id, user_id: info)
    app.dependency_overrides[get_current_user] = lambda: SimpleNamespace(id="u-1")
    app.dependency_overrides[get_db] = lambda: None
    try:
        client = TestClient(app)
        r = client.get(f"/api/v1/upload-jobs/{uuid.uuid4()}/data-gagal.xlsx")
        assert r.status_code == 200 and r.headers["content-type"].startswith(ep.XLSX_MEDIA_TYPE)
        info["rows"] = []
        assert client.get(f"/api/v1/upload-jobs/{uuid.uuid4()}/data-gagal.xlsx").status_code == 404
    finally:
        app.dependency_overrides.clear()
