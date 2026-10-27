"""
Unduh SEMUA data gagal hasil SIMPAN (TAHAP 2) satu upload sebagai Excel (.xlsx).

Sumber data: quarantine.rows dari batch simpan seluruh part grup upload (fitur/baris asli utuh, termasuk
GAGAL_DISIMPAN), tanpa batas jumlah. Data gagal TAHAP 1 (analisa) dibuat langsung di browser dari tabel
preview + table_preview (gis-dashboard/utils/failedRowsExcel.ts), jadi tidak disalin ke database.

Workbook: sheet "Data Gagal" (satu baris per data gagal, semua part) + sheet "Ringkasan".
"""
import io
import json
import re
from collections import Counter
from datetime import datetime
from typing import Any
from uuid import UUID
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.core.exceptions import bad_request, not_found
from app.services.upload.part_response import _label

WIB = ZoneInfo("Asia/Jakarta")
CELL_LIMIT = 32_767  # batas karakter satu sel Excel
TRUNC_NOTE = " … [DIPOTONG: melebihi batas 32.767 karakter per sel Excel]"
FORMULA_PREFIX = ("=", "+", "-", "@", "\t", "\r")
KEY_COLUMNS = (
    ("Estate", ("Est_ID", "EstID", "Est", "Estate", "EstID_2", "Estate_2", "kode_estate", "estate_code")),
    ("Afdeling", ("Afdeling", "Afdeling_2", "Divisi", "kode_afdeling", "division_code")),
    ("Blok", ("Blok", "Block", "kode_blok", "block_code")),
    ("OBJECTID", ("OBJECTID", "ObjectID", "objectid", "FID")),
)
FIXED_HEADERS = ["No", "No Data di File", "Kode Gagal", "Keterangan Kode", "Alasan Gagal"]


# ---------------------------------------------------------------- sel aman
def safe_cell(value: Any) -> Any:
    """Nilai sel: angka/bool apa adanya; teks dipotong ke batas Excel dan dicegah dari formula injection."""
    if value is None or isinstance(value, bool | int | float):
        return value
    if isinstance(value, dict | list):
        value = json.dumps(value, ensure_ascii=False, default=str)
    text_ = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(value))  # karakter kontrol ditolak openpyxl
    if text_.startswith(FORMULA_PREFIX):
        text_ = "'" + text_
    if len(text_) > CELL_LIMIT:
        text_ = text_[: CELL_LIMIT - len(TRUNC_NOTE)] + TRUNC_NOTE
    return text_


def geometry_text(geometry: Any) -> str | None:
    """Geometri -> WKT (fallback JSON bila tidak bisa dibaca shapely)."""
    if geometry in (None, {}, ""):
        return None
    try:
        from shapely.geometry import shape
        return shape(geometry).wkt
    except Exception:
        return json.dumps(geometry, ensure_ascii=False, default=str)


def split_source(source: Any) -> tuple[dict, Any, bool]:
    """(atribut, geometri, is_geojson) dari fitur GeoJSON atau baris Excel."""
    if isinstance(source, dict) and (source.get("type") == "Feature" or "geometry" in source):
        return dict(source.get("properties") or {}), source.get("geometry"), True
    if isinstance(source, dict):
        return dict(source), None, False
    return ({"data": source} if source is not None else {}), None, False


def key_values(props: dict) -> dict:
    lower = {str(k).lower(): v for k, v in props.items()}
    out = {}
    for label, keys in KEY_COLUMNS:
        out[label] = next((lower[k.lower()] for k in keys
                           if k.lower() in lower and lower[k.lower()] not in (None, "")), None)
    return out


# ---------------------------------------------------------------- kumpulkan
def _group_jobs(db: Session, job_id: UUID, user_id: UUID) -> list[dict]:
    job = db.execute(text("""
        SELECT id, kind, layer, upload_group_id, filename, bulan, tahun, finished_at, status
        FROM audit.upload_jobs WHERE id = :id AND user_id = :u
    """), {"id": str(job_id), "u": str(user_id)}).mappings().first()
    if job is None:
        raise not_found("Proses upload tidak ditemukan.")
    where = "id = :id" if job["upload_group_id"] is None else "upload_group_id = :g AND kind = :k"
    rows = db.execute(text(f"""
        SELECT id, kind, layer, filename, bulan, tahun, status, batch_id, result, part_index, feature_offset,
               finished_at, created_at
        FROM audit.upload_jobs WHERE user_id = :u AND {where}
        ORDER BY COALESCE(part_index, 1), created_at
    """), {"id": str(job_id), "u": str(user_id), "g": str(job["upload_group_id"]), "k": job["kind"]}).mappings().all()
    return [dict(r) for r in rows]


def _execute_rows(db: Session, jobs_: list[dict]) -> list[dict]:
    batch_ids = sorted({str(j["batch_id"]) for j in jobs_ if j["batch_id"]})
    if not batch_ids:
        return []
    data = db.execute(text("""
        SELECT source_id, reason_code, reason, data FROM quarantine.rows
        WHERE upload_batch_id = ANY(CAST(:b AS uuid[])) ORDER BY id
    """), {"b": batch_ids}).mappings().all()

    def _no(v):
        try:
            return int(v)
        except (TypeError, ValueError):
            return v
    return [{"no": _no(r["source_id"]), "kode": r["reason_code"], "alasan": r["reason"], "sumber": r["data"]}
            for r in data]


def _sort_key(row: dict):
    no = row.get("no")
    return (0, no, "") if isinstance(no, int) else (1, 0, str(no))


def collect(db: Session, job_id: UUID, user_id: UUID) -> dict:
    jobs_ = _group_jobs(db, job_id, user_id)
    first = jobs_[0]
    kind = first["kind"]
    if kind != "execute":
        raise bad_request("Data gagal tahap analisa diunduh langsung dari tabel preview di halaman upload.")
    rows, total = _execute_rows(db, jobs_), None
    finished = [j["finished_at"] for j in jobs_ if j["finished_at"]]
    return {
        "kind": kind,
        "layer": first["layer"],
        "filename": next((j["filename"] for j in jobs_ if j["filename"]), None),
        "bulan": first["bulan"],
        "tahun": first["tahun"],
        "waktu": max(finished) if finished else None,
        "total": total,
        "part_total": len(jobs_),
        "rows": sorted(rows, key=_sort_key),
    }


# ---------------------------------------------------------------- workbook
def _clean_filename(name: str | None) -> str:
    stem = re.sub(r"(_part\d+of\d+)?\.(geo)?json$|\.(xlsx|xls|csv)$", "", name or "", flags=re.I)
    stem = re.sub(r"[^A-Za-z0-9_-]+", "_", stem).strip("_")
    return stem[:60] or "file"


def export_filename(info: dict, now: datetime | None = None) -> str:
    now = (now or datetime.now(WIB)).astimezone(WIB)
    layer = re.sub(r"[^A-Za-z0-9_-]+", "_", info.get("layer") or "data")
    return f"data_gagal_{layer}_{_clean_filename(info.get('filename'))}_{now:%Y%m%d-%H%M}.xlsx"


def _fmt_time(value: Any) -> str:
    if isinstance(value, datetime):
        if value.tzinfo is None:
            value = value.replace(tzinfo=ZoneInfo("UTC"))
        return value.astimezone(WIB).strftime("%d-%m-%Y %H:%M WIB")
    return "-"


HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="B91C1C")
WRAP = Alignment(wrap_text=True, vertical="top")


def build_workbook(info: dict, now: datetime | None = None) -> bytes:
    rows = info["rows"]
    prepared, prop_keys, has_geom = [], [], False
    seen = set()
    for r in rows:
        props, geom, is_geo = split_source(r.get("sumber"))
        has_geom = has_geom or is_geo
        for k in props:
            if k not in seen:
                seen.add(k)
                prop_keys.append(k)
        prepared.append((r, props, geom))

    wb = Workbook()
    ws = wb.active
    ws.title = "Data Gagal"
    key_labels = [label for label, _ in KEY_COLUMNS]
    headers = FIXED_HEADERS + key_labels + [f"[file] {k}" for k in prop_keys] + (["Geometri (WKT)"] if has_geom else [])
    ws.append([safe_cell(h) for h in headers])
    for no, (r, props, geom) in enumerate(prepared, 1):
        keys = key_values(props)
        line = [no, r.get("no"), r.get("kode"), _label(r.get("kode") or "") if r.get("kode") else None, r.get("alasan")]
        line += [keys[label] for label in key_labels]
        line += [props.get(k) for k in prop_keys]
        if has_geom:
            line.append(geometry_text(geom))
        ws.append([safe_cell(v) for v in line])

    for cell in ws[1]:
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    widths = {1: 6, 2: 10, 3: 24, 4: 30, 5: 80}
    for idx in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(idx, 50 if headers[idx - 1] == "Geometri (WKT)" else 16)
    for row in ws.iter_rows(min_row=2, min_col=5, max_col=5):
        row[0].alignment = WRAP
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(len(prepared) + 1, 1)}"

    # ---------------- Ringkasan
    ss = wb.create_sheet("Ringkasan")
    stage = "Analisis (Tahap 1)" if info["kind"] == "analyze" else "Simpan (Tahap 2)"
    meta = [
        ("Nama File", info.get("filename") or "-"),
        ("Kategori", info.get("layer") or "-"),
        ("Periode", f"{int(info['bulan']):02d}-{info['tahun']}" if info.get("bulan") and info.get("tahun") else "-"),
        ("Tahap", stage),
        ("Waktu Proses", _fmt_time(info.get("waktu"))),
        ("Diunduh", _fmt_time(now or datetime.now(WIB))),
        ("Jumlah Part", info.get("part_total") or 1),
        ("Total Data di File", info["total"] if info.get("total") is not None else "-"),
        ("Total Data Gagal", len(rows)),
    ]
    for label, value in meta:
        ss.append([label, safe_cell(value)])
        ss.cell(row=ss.max_row, column=1).font = Font(bold=True)
    ss.append([])
    ss.append(["Kode Gagal", "Jumlah", "Keterangan"])
    head_row = ss.max_row
    for cell in ss[head_row]:
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
    for code, n in sorted(Counter(r.get("kode") or "-" for r in rows).items(), key=lambda x: (-x[1], x[0])):
        ss.append([safe_cell(code), n, safe_cell(_label(code))])
    ss.column_dimensions["A"].width = 28
    ss.column_dimensions["B"].width = 40
    ss.column_dimensions["C"].width = 60

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
